from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date

name = "Keelan O'Doherty"
course = "AEM 428"
prof = "Professor [Insert Name]"
title = "Special Project 1: Quasi-1D Nozzle Flow with Friction and Heat Transfer"

out_path = r"C:\Users\Keela\.openclaw\workspace\ODoherty_AEM428_SpecialProject1_Report_Draft.docx"

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

# MLA-style heading (top-left)
for line in [name, course, prof, date.today().strftime('%B %d, %Y')]:
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Title
p = doc.add_paragraph(title)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Intro
doc.add_paragraph(
    "This report presents a numerical analysis of compressible flow through a converging-diverging nozzle using Chapter 8 methods from Shapiro. "
    "The nozzle geometry is represented from supplied contour data and interpolated with a smooth polynomial-like spline (PCHIP). "
    "Flow property variation is computed from the throat location (x = 0) to the nozzle exit for four required scenarios: "
    "(a) isentropic flow, (b) friction only, (c) heat transfer only, and (d) friction plus heat transfer."
)

doc.add_paragraph("Input/Boundary Conditions", style=None)
doc.add_paragraph(
    "The simulation uses inlet total pressure p0 = 50 MPa (gage), inlet total temperature T0 = 3000 K, "
    "specific heat ratio k = 1.4, and gas constant R = 287 J/(kg·K). "
    "For non-isentropic cases, the friction coefficient is set to f = 0.005 and the heat-addition parameter is applied in the axial energy relation."
)

doc.add_paragraph("Numerical Method", style=None)
doc.add_paragraph(
    "The MATLAB solver evaluates area A(x), first derivative dA/dx, and second derivative terms needed for the differential Mach-number equation. "
    "Because the governing equation becomes indeterminate near Mach 1, the code starts slightly above unity at the throat and uses small axial increments (dx = 1e-4 m) with adaptive step halving near sensitive regions. "
    "Static temperature, velocity, density, and pressure are then recovered from the solved Mach-number distribution."
)

doc.add_paragraph("Case Definitions", style=None)
for item in [
    "Case A (Isentropic): f = 0, q = 0",
    "Case B (Friction only): f = 0.005, q = 0",
    "Case C (Heat transfer only): f = 0, q ≠ 0",
    "Case D (Friction + heat transfer): f = 0.005, q ≠ 0",
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph("Results and Discussion", style=None)
doc.add_paragraph(
    "The required outputs are three property-variation plots versus axial distance from throat (x = 0): Mach number, temperature, and pressure. "
    "Pressure is plotted on a semilogarithmic axis (semilogy) to make relative variation visible. "
    "Comparing all four cases on common axes highlights how viscous effects and heat addition alter acceleration, thermal state, and pressure decay along the nozzle."
)

# Figure placeholders
placeholders = [
    "[Insert Figure 1 here: Mach number vs x for Cases A–D]",
    "[Insert Figure 2 here: Temperature vs x for Cases A–D]",
    "[Insert Figure 3 here: Pressure vs x (semilogy) for Cases A–D]",
    "[Optional Figure 4 here: Nozzle geometry r(x) and throat location]",
]
for ph in placeholders:
    p = doc.add_paragraph(ph)
    p.runs[0].italic = True

# Interpretation section
doc.add_paragraph("Engineering Interpretation", style=None)
doc.add_paragraph(
    "The isentropic case provides the baseline acceleration and pressure-temperature trends expected from area variation alone. "
    "Adding friction generally reduces available acceleration and modifies pressure losses through viscous dissipation. "
    "Adding heat transfer shifts total-temperature behavior and can significantly change local Mach evolution. "
    "With both effects active, the combined trajectory reflects interaction between momentum and energy-source terms. "
    "Near the throat, very small step sizes are essential to maintain numerical stability and avoid nonphysical oscillations."
)

doc.add_paragraph("Code and Reproducibility", style=None)
doc.add_paragraph(
    "Primary MATLAB file used: ODoherty_AEM428_Project1_Solver2.m. "
    "The script generates all three required plots and includes distinct logic for the four assignment scenarios. "
    "Final submitted package should include this report, the MATLAB scripts used, and exported figure images with readable axes/legends."
)

doc.add_paragraph("References", style=None)
doc.add_paragraph("Shapiro, Ascher H. The Dynamics and Thermodynamics of Compressible Fluid Flow, Chapter 8.")

# Note placeholders user needs to edit
doc.add_paragraph("\nItems to update before submission:")
for item in [
    "Replace professor placeholder with actual instructor name.",
    "Insert exported MATLAB figures at the marked locations.",
    "Add any assignment-specific numerical observations from your final run outputs.",
]:
    doc.add_paragraph(item, style='List Bullet')

doc.save(out_path)
print(out_path)
